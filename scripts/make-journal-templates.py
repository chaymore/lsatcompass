"""Builds the two wrong answer journal templates in public/templates/.

Run after editing the text below:
    python3 -m venv .venv && .venv/bin/pip install python-docx openpyxl
    .venv/bin/python scripts/make-journal-templates.py
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = Path(__file__).resolve().parent.parent / "public" / "templates"
INK = RGBColor(0x0F, 0x0E, 0x0C)
MUTED = RGBColor(0x75, 0x6B, 0x5B)

FIELDS = [
    "Date",
    "The question (test, section, number)",
    "Question type",
    "My answer, and why I picked it",
    "Right answer, and why it's right",
    "What went wrong",
    "How I'll get it right next time",
    "Redid it? (date, right or wrong)",
]
EXAMPLE = [
    "Oct 12",
    "PT 141, LR section 1, question 14",
    "Flaw",
    "B. It sounded like the argument's conclusion, so it felt safe.",
    "D. It names the mistake: the argument treats a correlation as a cause.",
    "Fell for a trap. I picked an answer that described the argument instead of its flaw.",
    "Name the flaw in my own words before I read the answers.",
    "Oct 19, right",
]
# Shown as a hint in the document version, where there is room to paste the question itself.
DOC_FIELD_HINTS = {
    "The question (test, section, number)": "Paste the question here or add a screenshot, so you can redo it later.",
    "What went wrong": "Misread it, fell for a trap, missing a concept, or ran out of time?",
}
HOW_TO = [
    "Log every question you miss, and every one you guessed on, even if you got it right.",
    "Be honest about why your answer was tempting. That is the habit you are training out.",
    "Write one rule for next time. Ask AI to explain the question if you are not sure why the right answer is right.",
    "Redo each question a week later without looking at your notes. Review the journal every week and drill the question types that keep showing up.",
]


def shade(cell, hex_fill):
    props = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:val"), "clear")
    fill.set(qn("w:color"), "auto")
    fill.set(qn("w:fill"), hex_fill)
    props.append(fill)


def entry_table(doc, values):
    table = doc.add_table(rows=len(FIELDS), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row, (field, value) in zip(table.rows, zip(FIELDS, values)):
        row.cells[0].width = Inches(1.9)
        row.cells[1].width = Inches(4.6)
        shade(row.cells[0], "EDE7D9")
        label = row.cells[0].paragraphs[0].add_run(field)
        label.bold = True
        label.font.size = Pt(10)
        text = row.cells[1].paragraphs[0].add_run(value or DOC_FIELD_HINTS.get(field, ""))
        text.font.size = Pt(10)
        text.font.color.rgb = MUTED
        if not value:
            text.italic = True
            tall = field.startswith(("The question", "My answer", "Right answer", "How I'll"))
            row.height = Inches(1.4 if field.startswith("The question") else 0.6 if tall else 0.35)
    doc.add_paragraph()


def build_docx():
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(0.9)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK

    title = doc.add_paragraph()
    run = title.add_run("LSAT Wrong Answer Journal")
    run.font.name = "Georgia"
    run.font.size = Pt(24)
    sub = doc.add_paragraph()
    sub.add_run("From LSAT Compass, lsatcompass.com").font.color.rgb = MUTED

    doc.add_paragraph().add_run("How to use it").bold = True
    for line in HOW_TO:
        doc.add_paragraph(line, style="List Number")

    doc.add_paragraph().add_run("Example").bold = True
    entry_table(doc, EXAMPLE)

    doc.add_paragraph().add_run("Your entries").bold = True
    for _ in range(10):
        entry_table(doc, [""] * len(FIELDS))
    doc.add_paragraph("Copy a blank table to add more entries.").runs[0].font.color.rgb = MUTED
    doc.save(OUT / "lsat-wrong-answer-journal.docx")


def build_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Journal"
    widths = [11, 24, 16, 36, 36, 32, 36, 18]
    thin = Side(style="thin", color="DDD5C6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.append(FIELDS)
    ws.append(EXAMPLE)
    for col, width in enumerate(widths, start=1):
        letter = ws.cell(row=1, column=col).column_letter
        ws.column_dimensions[letter].width = width
        head = ws.cell(row=1, column=col)
        head.font = Font(bold=True, color="0F0E0C")
        head.fill = PatternFill("solid", fgColor="EDE7D9")
        head.alignment = Alignment(vertical="center", wrap_text=True)
        head.border = border
        for row in range(2, 102):
            cell = ws.cell(row=row, column=col)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            if row == 2:
                cell.font = Font(italic=True, color="756B5B")
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    tips = wb.create_sheet("How to use it")
    tips.column_dimensions["A"].width = 100
    tips.append(["LSAT Wrong Answer Journal, from LSAT Compass (lsatcompass.com)"])
    tips["A1"].font = Font(bold=True, size=14)
    for i, line in enumerate(HOW_TO, start=1):
        tips.append([f"{i}. {line}"])
    tips.append(["Row 2 on the Journal sheet is an example. Replace it with your own first entry."])
    wb.save(OUT / "lsat-wrong-answer-journal.xlsx")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    build_docx()
    build_xlsx()
    print("Wrote", *sorted(p.name for p in OUT.iterdir()))
